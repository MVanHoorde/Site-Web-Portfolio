#!/usr/bin/env python3
"""
gabarit_diapo.py — les briques communes à tous les diaporamas de projection PC.

Ce module ne connaît AUCUN chapitre. Il fournit le format de diapositive, la
charte, les composants (bandeau, encart, figure, formule, exercice, picto
fiche, checklist) et le **relevé des étapes d'animation**. Le contenu d'un
chapitre vit dans un fichier séparé, `diapo_<code>.py`.

POURQUOI CE FICHIER EXISTE — la chaîne qui avait produit le diaporama de
T3-C1 en août 2026 a vécu le temps d'une session, hors dépôt. Résultat : le
tableau des célérités de T3-C1 est faux et **rien ne peut le régénérer**, il
faut ouvrir PowerPoint. Ici la chaîne est versionnée : un diaporama se refait
en relançant un script.

QUATRE RÈGLES QUI COMMANDENT CE FICHIER
1. Aucune correction à l'écran. Les énoncés apparaissent, jamais les réponses.
2. Le picto ✎ se pose UNIQUEMENT là où la page en ligne porte un
   `<span class="a-noter">`. C'est ce qui garantit que l'écran et la fiche ne
   divergent pas.
3. Les figures viennent du site, extraites par `extraire_figures.mjs`. On ne
   redessine jamais dans PowerPoint ce qui existe déjà en ligne.
4. Polices sûres : Cambria, Calibri, Courier New. Space Grotesk et IBM Plex
   ne sont pas installées sur les postes du lycée ; la substitution serait
   incontrôlée en projection.
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# --------------------------------------------------------------- la charte

ENCRE = RGBColor(0x14, 0x1A, 0x26)
PAPIER = RGBColor(0xFD, 0xFC, 0xF8)
BETA = RGBColor(0x1D, 0x9A, 0xAA)      # teal   — définitions
GAMMA = RGBColor(0x4A, 0x3F, 0x9E)     # violet — propriétés, notations
ALPHA = RGBColor(0xD6, 0x40, 0x2B)     # rouge  — exercices, pièges
OCRE = RGBColor(0xB2, 0x8A, 0x1D)      # ocre   — méthode, marqueur fiche
GRIS = RGBColor(0x6B, 0x6F, 0x7A)
GRILLE = RGBColor(0xE8, 0xE4, 0xD8)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)

TITRES = "Cambria"
CORPS = "Calibri"
MONO = "Courier New"

# 16:9
LARGEUR, HAUTEUR = Emu(12192000), Emu(6858000)
CM = 360000          # EMU par centimètre
MARGE = 0.9          # cm, marge de sécurité aux bords


def cm(x):
    return Emu(int(x * CM))


# Fond pâle de chaque genre d'encart. Un contour noir sur chaque cadre
# alourdissait la diapositive (audit du 22/09 : « le cadre de définition est
# un peu lourd ») ; le fond teinté + le filet suffisent à délimiter.
TEINTES = {
    "definition": RGBColor(0xEA, 0xF5, 0xF6),
    "propriete": RGBColor(0xEF, 0xED, 0xF8),
    "notation": RGBColor(0xEF, 0xED, 0xF8),
    "exercice": RGBColor(0xFF, 0xFF, 0xFF),     # blanc + contour clair
    "methode": RGBColor(0xFB, 0xF5, 0xE3),
    "piege": RGBColor(0xFD, 0xEE, 0xEB),
    "exemple": RGBColor(0xEE, 0xF5, 0xEF),
}

GENRES = {
    "definition": (BETA, "DÉFINITION"),
    "propriete": (GAMMA, "PROPRIÉTÉ"),
    "notation": (GAMMA, "NOTATIONS"),
    "exercice": (ALPHA, "EXERCICE"),
    "methode": (OCRE, "MÉTHODE"),
    "piege": (ALPHA, "ATTENTION"),
    "exemple": (RGBColor(0x4C, 0x8A, 0x58), "EXEMPLE"),
}


# ------------------------------------------------------- relevé d'animation

class Sequence:
    """Retient quelles formes composent chaque étape d'animation.

    Le RANG est explicite parce que **l'ordre d'apparition n'est pas celui du
    code** : on construit souvent les encadrés avant les figures, alors que
    les figures doivent s'afficher d'abord — on les commente avec la classe
    avant d'institutionnaliser."""

    def __init__(self):
        self.par_diapo = {}     # index de diapo -> {rang: [shape_id, ...]}

    def fige(self, diapo, forme):
        """Forme présente dès l'arrivée : bandeau, pied, titre. Rang 0."""
        self.marque(diapo, 0, forme)

    def marque(self, diapo, rang, forme):
        if rang == 0:
            return              # rang 0 = visible d'emblée, pas d'animation
        d = self.par_diapo.setdefault(diapo, {})
        d.setdefault(rang, []).append(forme.shape_id)

    def etapes(self, diapo):
        """[[ids de l'étape 1], [ids de l'étape 2], ...] dans l'ordre."""
        d = self.par_diapo.get(diapo, {})
        return [d[r] for r in sorted(d)]

    def total(self):
        return sum(len(self.etapes(d)) for d in self.par_diapo)


# ------------------------------------------------------------- présentation

def presentation():
    p = Presentation()
    p.slide_width, p.slide_height = LARGEUR, HAUTEUR
    return p


def diapo(prs):
    """Une diapositive vierge, fond papier."""
    s = prs.slides.add_slide(prs.slide_layouts[6])   # disposition vide
    fond = s.background.fill
    fond.solid()
    fond.fore_color.rgb = PAPIER
    return s


# --------------------------------------------------------------- primitives

def _cadre(s, x, y, w, h, forme=MSO_SHAPE.RECTANGLE):
    return s.shapes.add_shape(forme, cm(x), cm(y), cm(w), cm(h))


def _sans_trait(shp):
    shp.line.fill.background()
    return shp


class Par:
    """Un paragraphe qui porte ses propres réglages : alignement, espace avant.

    Sert surtout aux ÉQUATIONS, qu'on centre au milieu d'un texte aligné à
    gauche — comme `.eq-ligne` sur le site. Le contenu est celui d'un
    paragraphe ordinaire : chaîne, ou liste de morceaux riches."""

    def __init__(self, contenu, align=None, avant=0.0, taille=None,
                 police=None, italique=None, couleur=None):
        self.contenu, self.align, self.avant = contenu, align, avant
        self.taille, self.police = taille, police
        self.italique, self.couleur = italique, couleur


def eq(contenu, taille=19):
    """Une équation centrée, en Cambria : la ligne `.eq-ligne` du site."""
    return Par(contenu, align=PP_ALIGN.CENTER, taille=taille, police=TITRES,
               avant=2)


def terme(texte, couleur=BETA):
    """Le mot DÉFINI, mis en évidence comme `.terme` sur le site.

    Le gras seul ne suffisait pas : dans un cadre de six lignes, l'œil ne
    trouvait pas « énergie massique » (audit du 22/09)."""
    return (texte, {"gras": True, "couleur": couleur})


def _texte(s, x, y, w, h, texte="", taille=16, police=CORPS, couleur=ENCRE,
           gras=False, italique=False, align=PP_ALIGN.LEFT,
           ancre=MSO_ANCHOR.TOP, interligne=1.0, espacement=0):
    """Zone de texte nue. `texte` peut être une liste de paragraphes."""
    zt = s.shapes.add_textbox(cm(x), cm(y), cm(w), cm(h))
    tf = zt.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancre
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    lignes = texte if isinstance(texte, (list, tuple)) else [texte]
    for i, ligne in enumerate(lignes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(ligne, Par):
            p.alignment = ligne.align if ligne.align is not None else align
            if ligne.avant:
                p.space_before = Pt(ligne.avant)
            p.line_spacing = interligne
            if espacement:
                p.space_after = Pt(espacement)
            _riche(p, ligne.contenu, ligne.taille or taille,
                   ligne.police or police, ligne.couleur or couleur, gras,
                   italique if ligne.italique is None else ligne.italique)
            continue
        p.alignment = align
        p.line_spacing = interligne
        if espacement:
            p.space_after = Pt(espacement)
        _riche(p, ligne, taille, police, couleur, gras, italique)
    return zt


def _riche(p, texte, taille, police, couleur, gras, italique):
    """Écrit `texte` dans le paragraphe `p`.

    `texte` peut être une chaîne, ou une liste de morceaux
    (texte, {options}) pour les indices, exposants et mises en couleur.
    Les indices passent par `subscript`, jamais par une espace : « U max »
    écrit avec une espace se lit comme deux mots au fond de la salle."""
    morceaux = texte if isinstance(texte, list) else [(texte, {})]
    for m in morceaux:
        t, o = m if isinstance(m, tuple) else (m, {})
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = o.get("police", police)
        f.size = Pt(o.get("taille", taille))
        f.bold = o.get("gras", gras)
        f.italic = o.get("italique", italique)
        f.color.rgb = o.get("couleur", couleur)
        if o.get("ind"):
            r.font._rPr.set("baseline", "-25000")
        if o.get("exp"):
            r.font._rPr.set("baseline", "30000")


def ind(base, indice):
    """Grandeur indicée : L_fusion, θ_changement d'état…"""
    return [(base, {"italique": True}), (indice, {"ind": True})]


def exp(base, exposant):
    return [(base, {}), (exposant, {"exp": True})]


# --------------------------------------------------------------- composants

def bandeau(s, seq, i, romain, partie, sous=""):
    """Bandeau de tête : chiffre romain, partie, sous-partie. Visible d'emblée."""
    barre = _sans_trait(_cadre(s, 0, 0, 33.87, 1.75))
    barre.fill.solid()
    barre.fill.fore_color.rgb = ENCRE
    barre.text_frame.text = ""
    seq.fige(i, barre)

    # les trois repères du « spectre » du site, groupés à droite : posés
    # à gauche, ils venaient buter contre le titre de la partie.
    for col, gauche in ((GAMMA, 30.9), (BETA, 31.4), (ALPHA, 31.9)):
        t = _sans_trait(_cadre(s, gauche, 0, 0.12, 1.75))
        t.fill.solid()
        t.fill.fore_color.rgb = col
        seq.fige(i, t)

    z = _texte(s, MARGE, 0.28, 3.0, 1.2, romain, taille=26, police=TITRES,
               couleur=PAPIER, gras=True, ancre=MSO_ANCHOR.MIDDLE)
    seq.fige(i, z)
    z = _texte(s, MARGE + 2.6, 0.30, 24.0, 0.7, partie, taille=17,
               police=TITRES, couleur=PAPIER, gras=True)
    seq.fige(i, z)
    if sous:
        z = _texte(s, MARGE + 2.6, 1.02, 24.0, 0.6, sous, taille=11,
                   police=MONO, couleur=RGBColor(0xC8, 0xC4, 0xB8))
        seq.fige(i, z)


def pied(s, seq, i, chapitre, numero):
    """Pied de page. Visible d'emblée."""
    filet = _sans_trait(_cadre(s, MARGE, 17.4, 33.87 - 2 * MARGE, 0.04))
    filet.fill.solid()
    filet.fill.fore_color.rgb = GRILLE
    seq.fige(i, filet)
    z = _texte(s, MARGE, 17.6, 24.0, 0.6, chapitre, taille=10, police=MONO,
               couleur=GRIS)
    seq.fige(i, z)
    z = _texte(s, 33.87 - MARGE - 4.0, 17.6, 4.0, 0.6, str(numero), taille=10,
               police=MONO, couleur=GRIS, align=PP_ALIGN.RIGHT)
    seq.fige(i, z)


def picto_fiche(s, seq, i, rang, x, y, libelle=True):
    """Pastille ocre ✎ + « SUR LA FICHE ». (x, y) = coin de la PASTILLE.

    🔴 Se pose UNIQUEMENT là où la page en ligne porte un `a-noter`.

    Le libellé est posé à GAUCHE de la pastille et aligné à droite : placé à
    sa droite, il sortait du cadre de l'encart dès que la pastille était
    contre le bord — le texte était coupé sur les huit premières diapos."""
    if libelle:
        z = _texte(s, x - 4.15, y + 0.16, 4.0, 0.5, "SUR LA FICHE", taille=9,
                   police=MONO, couleur=OCRE, gras=True, align=PP_ALIGN.RIGHT)
        seq.marque(i, rang, z)
    p = _cadre(s, x, y, 0.72, 0.72, MSO_SHAPE.OVAL)
    p.fill.solid()
    p.fill.fore_color.rgb = OCRE
    p.line.color.rgb = OCRE
    tf = p.text_frame
    tf.text = "✎"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].runs[0]
    r.font.size, r.font.name, r.font.color.rgb = Pt(13), CORPS, PAPIER
    seq.marque(i, rang, p)
    return p


PT_CM = 0.03528          # 1 point typographique en centimètres

# Largeur moyenne d'un caractère, en cadratins (em). Mesurée sur les trois
# polices sûres du diaporama — Courier New est monospace à 0,60 em, d'où des
# blocs « Données : … » bien plus larges que le même texte en Calibri.
AVANCE = {"Calibri": 0.48, "Cambria": 0.50, "Courier New": 0.60}


def _avance(police, pt):
    """Largeur moyenne d'un caractère, en cm, pour cette police à cette taille."""
    return AVANCE.get(police, 0.50) * pt * PT_CM


def bas(forme, jeu=0.55):
    """Le bas d'une forme, en cm, plus un jeu — pour empiler le bloc suivant.

    POURQUOI — les cadres ont désormais une hauteur calculée à partir de leur
    texte. Une ordonnée codée en dur redevient fausse dès qu'on change une
    phrase : sur la diapositive des équations à équilibrer, l'écart entre les
    deux exercices avait atteint 5 cm. On empile relativement, pas
    absolument."""
    return round((forme.top + forme.height) / CM + jeu, 2)


def hauteur_encart(corps, w, taille=15, police=CORPS, etiquette_sur_2=False):
    """Estime la hauteur qu'il faut à un encart pour son texte.

    POURQUOI — sans cela, chaque cadre demande un réglage à la main, et une
    boîte deux fois trop haute pour trois lignes se voit à la projection.

    🔴 **Chaque paragraphe est mesuré à SA taille.** L'estimation qui
    appliquait 15 pt à tout le cadre faisait sortir les quatre équations de
    l'exercice 11, écrites en 22 pt, hors de leur cadre — 1,3 cm dehors, et
    la quatrième équation flottait sur le fond. Une taille de run, ça change
    la largeur des caractères ET la hauteur de ligne.

    Cette estimation ne se vérifie pas toute seule : `mesurer_cadres.ps1`
    demande à PowerPoint la hauteur réelle de chaque texte.
    """
    tete = 1.35 if etiquette_sur_2 else 0.95
    # 0,45 cm de garde : l'estimation reste une estimation, et un cadre un peu
    # trop grand se voit moins qu'un texte qui en sort.
    return round(tete + hauteur_texte(corps, w - 1.05, taille, police) + 0.45, 2)


def hauteur_texte(corps, utile, taille=15, police=CORPS):
    """Hauteur, en cm, d'une liste de paragraphes dans une colonne `utile`."""
    haut = 0.0
    for p in corps:
        if isinstance(p, Par):
            haut += p.avant * PT_CM
            pt_par, pol_par = p.taille or taille, p.police or police
            p = p.contenu
        else:
            pt_par, pol_par = taille, police
        morceaux = ([(p, {})] if isinstance(p, str)
                    else [m if isinstance(m, tuple) else (m, {}) for m in p])
        # largeur du paragraphe = somme des largeurs de ses runs, chacun à SA
        # police et SA taille : un « Données : … » en Courier 12 n'occupe pas
        # la même place que le même texte en Calibri 15.
        # une espace ne pèse qu'une demi-lettre : les équations, aérées
        # d'espaces multiples, étaient estimées sur deux lignes et laissaient
        # un creux dans leur cadre
        largeur = sum((len(t) - 0.5 * t.count(" "))
                      * _avance(o.get("police", pol_par),
                                o.get("taille", pt_par))
                      for t, o in morceaux)
        pt = max(o.get("taille", pt_par) for _, o in morceaux)
        lignes = max(1, -(-largeur // utile))
        haut += lignes * PT_CM * pt * 1.22 + 0.21
    return haut


def _habiller(cadre, genre):
    """Fond pâle du genre ; contour clair pour l'exercice seul (fond blanc)."""
    cadre.fill.solid()
    cadre.fill.fore_color.rgb = TEINTES.get(genre, BLANC)
    if genre == "exercice":
        cadre.line.color.rgb = RGBColor(0xD8, 0xD3, 0xC4)
        cadre.line.width = Pt(0.75)
    else:
        cadre.line.fill.background()
    cadre.shadow.inherit = False
    cadre.text_frame.text = ""


def encart(s, seq, i, rang, genre, x, y, w, h, corps, etiquette=None,
           fiche=False):
    """Cadre à filet gauche coloré : définition, propriété, exercice, méthode.

    `corps` = liste de paragraphes (chaîne, ou liste de morceaux riches).
    `valign` reste en HAUT : centré, le texte creuse un vide sous l'étiquette."""
    couleur, defaut = GENRES[genre]
    if h is None:                       # hauteur ajustée au texte
        h = hauteur_encart(corps, w)
    cadre = _cadre(s, x, y, w, h)
    _habiller(cadre, genre)
    seq.marque(i, rang, cadre)

    filet = _sans_trait(_cadre(s, x, y, 0.13, h))
    filet.fill.solid()
    filet.fill.fore_color.rgb = couleur
    seq.marque(i, rang, filet)

    z = _texte(s, x + 0.45, y + 0.28, w - 1.0, 0.5,
               etiquette or defaut, taille=10, police=MONO, couleur=couleur,
               gras=True)
    seq.marque(i, rang, z)

    z = _texte(s, x + 0.45, y + 0.95, w - 0.9, h - 1.25, corps, taille=15,
               interligne=1.15, espacement=6)
    seq.marque(i, rang, z)

    if fiche:
        # étiquette longue dans un cadre étroit : « SUR LA FICHE » la percutait
        # (diapositive 9 de T1-C2). La pastille seule suffit alors — sa légende
        # est posée dès la page de titre.
        place = w - 1.0 - len(etiquette or defaut) * _avance(MONO, 10)
        picto_fiche(s, seq, i, rang, x + w - 1.05, y + 0.2,
                    libelle=place > 4.6)
    return cadre


def _mesures_exercice(w, contexte=(), equations=(), fig=None, donnees=None,
                      question=None, colonnes=None):
    """Les hauteurs de chaque partie d'un exercice, sans rien dessiner."""
    # contexte : une chaîne, un Par, ou une LISTE de paragraphes (un
    # paragraphe riche s'écrit donc [[(…), (…)]])
    if not contexte:
        contexte = []
    elif isinstance(contexte, (str, Par)):
        contexte = [contexte]
    haut_ctx = list(contexte) + [eq(e) for e in equations]
    utile = w - 1.05
    m = {"haut_ctx": haut_ctx,
         "h_ctx": hauteur_texte(haut_ctx, utile) if haut_ctx else 0.0,
         "h_col": 0.0}
    if colonnes:
        ncol, items = colonnes
        m["par_col"] = -(-len(items) // ncol)
        m["largeur_col"] = (utile - 0.5 * (ncol - 1)) / ncol
        m["h_col"] = max(hauteur_texte(items[k * m["par_col"]:(k + 1) * m["par_col"]],
                                       m["largeur_col"], taille=14)
                         for k in range(ncol)) + 0.1
    m["h_fig"] = (fig[2] * fig[1] + 0.35) if fig else 0.0
    bas_txt = ([Par(donnees, taille=12, police=MONO, couleur=GRIS)]
               if donnees else [])
    if question:
        bas_txt.append(Par(question, italique=True))
    m["bas_txt"] = bas_txt
    m["h_bas"] = hauteur_texte(bas_txt, utile) if bas_txt else 0.0
    # 0,25 cm entre ce qu'on DONNE et ce qu'on DEMANDE : sans lui, la
    # question collait à la dernière équation
    m["souffle"] = 0.25 if (bas_txt and (haut_ctx or colonnes or fig)) else 0.0
    m["h"] = round(0.95 + m["h_ctx"] + m["h_col"] + m["h_fig"] + m["souffle"]
                   + m["h_bas"] + 0.4, 2)
    return m


def hauteur_exercice(w, **k):
    """Hauteur qu'occupera un exercice — pour aligner deux exercices posés
    face à face sur la même hauteur (audit du 22/09 : exercices 4 et 5)."""
    return _mesures_exercice(w, **k)["h"]


def exercice(s, seq, i, rang, x, y, w, num, titre, contexte=(), equations=(),
             fig=None, donnees=None, question=None, colonnes=None, h=None):
    """Un exercice, structuré comme sur le site.

    contexte   — paragraphe(s) en romain : la situation
    colonnes   — liste d'items répartis en N colonnes (ex. 1, neuf situations)
    equations  — équations CENTRÉES, en Cambria : la ligne `.eq-ligne`
    fig        — (chemin, ratio h/w, largeur cm) : la figure vit DANS le
                 cadre (R7) — posée à côté, elle se lisait comme une
                 illustration du cours
    donnees    — en Courier, gris : la ligne `.donnees`
    question   — en italique : la ligne `.question`

    POURQUOI — l'énoncé entier était écrit en un seul paragraphe italique :
    contexte, équation et question se soudaient en un bloc (audit du 22/09).
    Retourne le cadre, pour empiler le bloc suivant avec `bas()`."""
    m = _mesures_exercice(w, contexte, equations, fig, donnees, question,
                          colonnes)
    haut_ctx, bas_txt = m["haut_ctx"], m["bas_txt"]
    h_ctx, h_col, h_fig, h_bas = m["h_ctx"], m["h_col"], m["h_fig"], m["h_bas"]
    if colonnes:
        ncol, items = colonnes
        par_col, largeur_col = m["par_col"], m["largeur_col"]
    y_txt = y + 0.95
    if h is None:
        h = m["h"]

    cadre = _cadre(s, x, y, w, h)
    _habiller(cadre, "exercice")
    seq.marque(i, rang, cadre)
    filet = _sans_trait(_cadre(s, x, y, 0.13, h))
    filet.fill.solid()
    filet.fill.fore_color.rgb = ALPHA
    seq.marque(i, rang, filet)
    z = _texte(s, x + 0.45, y + 0.28, w - 1.0, 0.5,
               f"EXERCICE {num} — {titre.upper()}", taille=10, police=MONO,
               couleur=ALPHA, gras=True)
    seq.marque(i, rang, z)

    cur = y_txt
    if haut_ctx:
        z = _texte(s, x + 0.45, cur, w - 0.9, h_ctx, haut_ctx, taille=15,
                   interligne=1.15, espacement=6)
        seq.marque(i, rang, z)
        cur += h_ctx
    if colonnes:
        for k in range(ncol):
            bloc = items[k * par_col:(k + 1) * par_col]
            if not bloc:
                continue
            z = _texte(s, x + 0.45 + k * (largeur_col + 0.5), cur,
                       largeur_col, h_col, bloc, taille=14, interligne=1.1,
                       espacement=5)
            seq.marque(i, rang, z)
        cur += h_col
    if fig:
        chemin, ratio, fw = fig
        img = s.shapes.add_picture(str(chemin), cm(x + (w - fw) / 2),
                                   cm(cur + 0.1), width=cm(fw),
                                   height=cm(fw * ratio))
        seq.marque(i, rang, img)
        cur += h_fig
    if bas_txt:
        cur += m["souffle"]
        z = _texte(s, x + 0.45, cur, w - 0.9, h_bas, bas_txt, taille=15,
                   interligne=1.15, espacement=6)
        seq.marque(i, rang, z)
    return cadre


def figure(s, seq, i, rang, chemin, x, y, w, legende_txt=None, h=None):
    """Une figure du site, avec sa légende. Les figures passent en PREMIER
    dans l'ordre d'animation : on les commente avant d'institutionnaliser."""
    img = s.shapes.add_picture(str(chemin), cm(x), cm(y), width=cm(w),
                               height=cm(h) if h else None)
    seq.marque(i, rang, img)
    if legende_txt:
        bas = img.top + img.height + Emu(int(0.18 * CM))
        z = s.shapes.add_textbox(cm(x), bas, cm(w), cm(1.2))
        tf = z.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        _riche(tf.paragraphs[0], legende_txt, 10, MONO, GRIS, False, False)
        tf.paragraphs[0].line_spacing = 1.1
        seq.marque(i, rang, z)
    return img


def fraction(s, seq, i, rang, gauche, haut, bas, x, y, w=5.0, couleur=ENCRE):
    """Fraction numérateur SUR dénominateur — jamais de barre oblique.

    Trois objets PowerPoint natifs : rien n'est aplati en image, on peut
    cliquer le numérateur et le changer. Le membre de gauche est aligné à
    DROITE pour que le « = » bute contre la fraction.

    `couleur` : PAPIER pour poser la fraction DANS le panneau sombre de
    `formule()` — à l'encre, elle y était invisible (T1-C1, ρ = m/V)."""
    zg = _texte(s, x - 5.0, y + 0.55, 4.8, 0.9, gauche, taille=22,
                police=TITRES, couleur=couleur, italique=True,
                align=PP_ALIGN.RIGHT)
    seq.marque(i, rang, zg)
    zn = _texte(s, x, y, w, 0.85, haut, taille=20, police=TITRES,
                couleur=couleur, italique=True, align=PP_ALIGN.CENTER)
    seq.marque(i, rang, zn)
    trait = _sans_trait(_cadre(s, x, y + 0.92, w, 0.05))
    trait.fill.solid()
    trait.fill.fore_color.rgb = couleur
    seq.marque(i, rang, trait)
    zd = _texte(s, x, y + 1.05, w, 0.85, bas, taille=20, police=TITRES,
                couleur=couleur, italique=True, align=PP_ALIGN.CENTER)
    seq.marque(i, rang, zd)


def formule(s, seq, i, rang, expression, grandeurs, x, y, w=13.0, h=3.4):
    """Panneau encre + formule en réserve claire + « Grandeurs & unités »."""
    panneau = _sans_trait(_cadre(s, x, y, w, h))
    panneau.fill.solid()
    panneau.fill.fore_color.rgb = ENCRE
    panneau.text_frame.text = ""
    seq.marque(i, rang, panneau)

    z = _texte(s, x + 0.6, y + 0.5, w * 0.48, h - 1.0, expression, taille=26,
               police=TITRES, couleur=PAPIER, italique=True,
               align=PP_ALIGN.CENTER, ancre=MSO_ANCHOR.MIDDLE)
    seq.marque(i, rang, z)

    sep = _sans_trait(_cadre(s, x + w * 0.52, y + 0.5, 0.03, h - 1.0))
    sep.fill.solid()
    sep.fill.fore_color.rgb = GRIS
    seq.marque(i, rang, sep)

    z = _texte(s, x + w * 0.57, y + 0.45, w * 0.40, 0.5,
               "GRANDEURS & UNITÉS", taille=9, police=MONO,
               couleur=RGBColor(0x9A, 0x8F, 0x6E), gras=True)
    seq.marque(i, rang, z)
    z = _texte(s, x + w * 0.57, y + 1.05, w * 0.40, h - 1.5, grandeurs,
               taille=13, couleur=PAPIER, interligne=1.2, espacement=4)
    seq.marque(i, rang, z)


def titre_partie(s, seq, i, romain, titre, sous_parties):
    """Diapositive d'intercalaire : le grand chiffre romain et le sommaire."""
    bloc = _sans_trait(_cadre(s, 0, 0, 33.87, 19.05))
    bloc.fill.solid()
    bloc.fill.fore_color.rgb = ENCRE
    bloc.text_frame.text = ""
    seq.fige(i, bloc)
    # Le chiffre romain est AU-DESSUS du titre, pas derrière : posé en
    # filigrane, il passait sous le titre et sous les premières sous-parties,
    # et l'empilement se lisait comme une collision.
    z = _texte(s, 2.4, 3.0, 12.0, 4.2, romain, taille=104, police=TITRES,
               couleur=RGBColor(0x3A, 0x44, 0x5C), gras=True)
    seq.fige(i, z)
    z = _texte(s, 2.4, 7.5, 28.0, 2.0, titre, taille=36, police=TITRES,
               couleur=PAPIER, gras=True)
    seq.fige(i, z)
    filet = _sans_trait(_cadre(s, 2.4, 9.3, 5.0, 0.06))
    filet.fill.solid()
    filet.fill.fore_color.rgb = ALPHA
    seq.fige(i, filet)
    # 🔴 TOUTES les sous-parties au MÊME rang : elles apparaissent d'un seul
    # clic. Une par clic, c'était sept clics pour le seul sommaire de la
    # partie III — on passe le temps de la projection à cliquer.
    for k, sp in enumerate(sous_parties):
        z = _texte(s, 2.6, 10.1 + k * 0.95, 28.0, 0.8, sp, taille=15,
                   police=MONO, couleur=RGBColor(0x9A, 0xA2, 0xB2))
        seq.marque(i, 1, z)


def checklist(s, seq, i, titre, items, kahoots=()):
    """« Pour le DS, je sais » — les compétences se dévoilent une par une,
    chacune cochée à l'oral avant d'apparaître.

    `kahoots` : [(libellé, url), …] — le ou les Kahoot bilan du chapitre, ceux
    de la checklist en ligne. 🔴 OBLIGATOIRES dès que la page en porte
    (demande de Loïc, 22/09/2026) : `controler.py` échoue s'il en manque un.
    Ils apparaissent au DERNIER clic, une fois toutes les compétences passées.
    Le lien est posé sur le RUN, jamais sur la zone de texte : posé sur la
    zone, il écrivait un r:id vide et PowerPoint déclarait le fichier corrompu."""
    for k, (libelle, url) in enumerate(kahoots):
        z = _texte(s, MARGE + 0.2 + k * 16.0, 2.6, 15.6, 0.8,
                   [[("► KAHOOT — ", {"gras": True}), (libelle, {})]],
                   taille=13, police=MONO, couleur=OCRE)
        for r in z.text_frame.paragraphs[0].runs:
            r.hyperlink.address = url
        seq.marque(i, len(items) + 1, z)
    z = _texte(s, MARGE + 0.2, 2.6, 30.0, 1.0, titre, taille=24, police=TITRES,
               gras=True)
    seq.fige(i, z)
    # le pas se resserre avec le nombre de compétences : à 1,55 cm fixe, la
    # neuvième de T1-C1 tombait sous le pied de page (bas à 17,9 cm).
    pas = min(1.55, 12.8 / max(1, len(items)))
    for k, it in enumerate(items):
        y = 4.1 + k * pas
        case = _cadre(s, MARGE + 0.4, y, 0.62, 0.62)
        case.fill.solid()
        case.fill.fore_color.rgb = BLANC
        case.line.color.rgb = ENCRE
        case.text_frame.text = ""
        seq.marque(i, k + 1, case)
        z = _texte(s, MARGE + 1.4, y - 0.05, 29.0, 1.4, it, taille=14,
                   interligne=1.1)
        seq.marque(i, k + 1, z)


def diapo_titre(prs, seq, theme, chapitre, sous_titre, annee, fond=None,
                logo=None):
    """Page de titre. Le logo n'apparaît QUE là, sur une plaque claire :
    sur fond sombre, son texte bleu nuit disparaîtrait."""
    s = diapo(prs)
    i = 0
    bloc = _sans_trait(_cadre(s, 0, 0, 33.87, 19.05))
    bloc.fill.solid()
    bloc.fill.fore_color.rgb = ENCRE
    bloc.text_frame.text = ""
    if fond and Path(fond).exists():
        s.shapes.add_picture(str(fond), cm(0), cm(4.2), width=cm(33.87))
    _texte(s, 2.4, 5.4, 26.0, 0.8, theme, taille=13, police=MONO,
           couleur=BETA, gras=True)
    # le titre tient sur deux lignes : la boîte fait 3,4 cm, sans quoi le
    # sous-titre lui passe dessus.
    _texte(s, 2.4, 6.2, 27.0, 3.4, chapitre, taille=40, police=TITRES,
           couleur=PAPIER, gras=True, interligne=1.08)
    _texte(s, 2.4, 10.1, 28.0, 0.9, sous_titre, taille=16, police=CORPS,
           couleur=RGBColor(0x9A, 0xA2, 0xB2), italique=True)
    filet = _sans_trait(_cadre(s, 2.4, 11.5, 6.0, 0.06))
    filet.fill.solid()
    filet.fill.fore_color.rgb = ALPHA
    _texte(s, 2.4, 12.1, 26.0, 1.6,
           ["M. VAN HOORDE · LYCÉE ISAAC DE L'ÉTOILE",
            f"Seconde générale · année {annee}"],
           taille=12, police=MONO, couleur=RGBColor(0x9A, 0xA2, 0xB2),
           espacement=4)

    if logo and Path(logo).exists():
        # la plaque doit CONTENIR le logo : à 5,6 cm de large il fait
        # 3,15 cm de haut et débordait d'une plaque de 2,9 cm.
        plaque = _sans_trait(_cadre(s, 26.3, 2.0, 5.7, 3.3))
        plaque.fill.solid()
        plaque.fill.fore_color.rgb = PAPIER
        plaque.text_frame.text = ""
        s.shapes.add_picture(str(logo), cm(26.9), cm(2.35), width=cm(4.5))

    # la légende du picto ✎ est posée DÈS la page de titre : l'élève doit
    # savoir à quoi il sert avant d'en croiser un.
    p = _cadre(s, 2.4, 15.6, 0.72, 0.72, MSO_SHAPE.OVAL)
    p.fill.solid()
    p.fill.fore_color.rgb = OCRE
    p.line.color.rgb = OCRE
    tf = p.text_frame
    tf.text = "✎"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].runs[0]
    r.font.size, r.font.name, r.font.color.rgb = Pt(13), CORPS, PAPIER
    _texte(s, 3.55, 15.75, 24.0, 0.7,
           "ce picto marque ce qui est à écrire sur ta fiche de cours",
           taille=12, police=MONO, couleur=OCRE)
    return s
